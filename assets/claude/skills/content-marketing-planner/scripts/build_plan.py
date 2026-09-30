"""
build_plan.py
=============
Xây file Excel "Plan Content Marketing" theo đúng bố cục & format mẫu
(Nhà Hàng Hải Sản Biển Đông). Giữ nguyên 100% form:

- Font: Be Vietnam Pro
- Màu header chính: #1C4587 (xanh đậm), #0B5394 (xanh đậm TikTok)
- Màu header phụ: #1155CC (xanh trung bình)
- Màu nền content: #C9DAF8 (xanh nhạt), #CFE2F3 (xanh nhạt TikTok)
- Chữ header: trắng, bold, center/center
- Content: wrap_text=True, vertical=center
- Tất cả cell có border thin

Dùng:
    from build_plan import build_plan
    build_plan(data, "output.xlsx")

Cấu trúc data: xem docstring của build_plan() hoặc file examples/*.json
"""

from __future__ import annotations

import json
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ============================================================
# HẰNG SỐ ĐỊNH DẠNG (giữ nguyên 100% theo mẫu Biển Đông)
# ============================================================

FONT_NAME = "Be Vietnam Pro"

COLOR = {
    "dark_blue_1":   "FF1C4587",   # header chính (FB, CDKH, ĐHC, mapping sheets)
    "dark_blue_2":   "FF0B5394",   # header chính (TikTok)
    "medium_blue":   "FF1155CC",   # header phụ (CDKH segments)
    "medium_blue_2": "FF6D9EEB",   # ô tên Pillar trong FB pillar table
    "medium_blue_3": "FF3C78D8",   # sub-header (Ảnh thật/Ảnh thiết kế/Reel)
    "light_blue_1":  "FFC9DAF8",   # nền content (FB theme)
    "light_blue_2":  "FFCFE2F3",   # nền content (TikTok theme)
    "white":         "FFFFFFFF",
}

THIN = Side(style="thin", color="FF000000")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ============================================================
# HELPERS STYLE
# ============================================================

def _font(bold: bool = False, color: str = "FF000000", size: Optional[float] = 11.0) -> Font:
    return Font(name=FONT_NAME, bold=bold, color=color, size=size)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=hex_color, end_color=hex_color)


def _align(h: Optional[str] = None, v: str = "center", wrap: bool = True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _style_header(cell, fill_hex: str = COLOR["dark_blue_1"], h: str = "center") -> None:
    """Ô header (dark blue bg, white bold)."""
    cell.font = _font(bold=True, color=COLOR["white"])
    cell.fill = _fill(fill_hex)
    cell.alignment = _align(h=h, v="center", wrap=True)
    cell.border = BORDER_ALL


def _style_content(cell, fill_hex: str = COLOR["light_blue_1"], h: Optional[str] = None) -> None:
    """Ô nội dung (light blue bg, black text)."""
    cell.font = _font(bold=False, color="FF000000")
    cell.fill = _fill(fill_hex)
    cell.alignment = _align(h=h, v="center", wrap=True)
    cell.border = BORDER_ALL


def _style_plain(cell, h: Optional[str] = None) -> None:
    """Ô không có fill (dùng cho ô trắng trong calendar)."""
    cell.font = _font(bold=False, color="FF000000")
    cell.alignment = _align(h=h, v="center", wrap=True)
    cell.border = BORDER_ALL


def _merge_consecutive_same_values(ws: Worksheet, row: int, pillars: list, key: str) -> None:
    """Merge các ô liên tiếp có cùng value trên 1 dòng (dùng cho audience row)."""
    if not pillars:
        return
    start = 0
    current_val = pillars[0].get(key, "")
    for i in range(1, len(pillars) + 1):
        if i == len(pillars) or pillars[i].get(key, "") != current_val:
            if i - start > 1:  # Có >=2 ô liên tiếp giống nhau -> merge
                ws.merge_cells(
                    start_row=row, start_column=2 + start,
                    end_row=row, end_column=2 + i - 1,
                )
            if i < len(pillars):
                start = i
                current_val = pillars[i].get(key, "")


# ============================================================
# SHEET 1: CHÂN DUNG KHÁCH HÀNG
# ============================================================

def build_customer_portrait_sheet(wb: Workbook, data: Dict[str, Any]) -> None:
    """
    data["company_name"]: str - Tên doanh nghiệp (dùng cho title)
    data["customer_portrait"]:
        {
            "segments": ["Nhóm khách A", "Nhóm khách B", ...],  # 2-5 cột
            "dimensions": [
                {"label": "1. Nhân khẩu học", "values": ["text A", "text B", ...]},
                {"label": "3. Hành vi tiêu dùng", "values": [...]},
                ...
            ]
        }
    """
    ws = wb.create_sheet("CHÂN DUNG KHÁCH HÀNG")
    portrait = data["customer_portrait"]
    segments = portrait["segments"]
    dims = portrait["dimensions"]

    n_seg = len(segments)
    # Cột A = label, B..(B+n_seg-1) = segments
    total_cols = 1 + n_seg
    last_col_letter = get_column_letter(total_cols)

    # --- Row 1: Title ---
    title = f"CHÂN DUNG KHÁCH HÀNG {data['company_name'].upper()}"
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(f"A1:{last_col_letter}1")
    for c in range(1, total_cols + 1):
        _style_header(ws.cell(row=1, column=c))

    # --- Row 2: Headers ---
    ws.cell(row=2, column=1, value="Yếu tố")
    _style_header(ws.cell(row=2, column=1), COLOR["dark_blue_1"])
    for i, seg in enumerate(segments):
        cell = ws.cell(row=2, column=2 + i, value=seg)
        _style_header(cell, COLOR["medium_blue"])

    # --- Rows 3+: Dimensions ---
    for r_idx, dim in enumerate(dims, start=3):
        # Column A: label (dark blue bg, white bold)
        label_cell = ws.cell(row=r_idx, column=1, value=dim["label"])
        _style_header(label_cell, COLOR["dark_blue_1"], h=None)

        # Columns B+: segment values (light blue bg)
        for i, val in enumerate(dim["values"]):
            cell = ws.cell(row=r_idx, column=2 + i, value=val)
            _style_content(cell, COLOR["light_blue_1"])

    # --- Column widths (giữ nguyên mẫu) ---
    ws.column_dimensions["A"].width = 20.38
    ws.column_dimensions["B"].width = 33.63
    for i in range(n_seg - 1):
        # Các cột C, D, E... width 13 như mẫu
        ws.column_dimensions[get_column_letter(3 + i)].width = 13.0


# ============================================================
# SHEET 2: ĐỊNH HƯỚNG CHUNG
# ============================================================

def build_general_direction_sheet(wb: Workbook, data: Dict[str, Any]) -> None:
    """
    data["general_direction"]: list of {"label": "...", "value": "..."}
        Thường gồm 5 mục:
        - Mục tiêu chính
        - Định vị hình ảnh
        - Gợi ý định vị
        - Gợi ý BIO
        - Mục tiêu cụ thể
    """
    ws = wb.create_sheet("ĐỊNH HƯỚNG CHUNG")
    items = data["general_direction"]

    # Row 1 để trống như mẫu, bắt đầu từ row 2
    for i, item in enumerate(items, start=2):
        label_cell = ws.cell(row=i, column=2, value=item["label"])
        _style_header(label_cell, COLOR["dark_blue_1"], h="center")

        content_cell = ws.cell(row=i, column=3, value=item["value"])
        _style_content(content_cell, COLOR["light_blue_1"], h="left")

    ws.column_dimensions["B"].width = 16.38
    ws.column_dimensions["C"].width = 75.25


# ============================================================
# SHEET 3+: PLATFORM STRATEGY (e.g., FACEBOOK, TIKTOK, ...)
# ============================================================

def build_platform_sheet(wb: Workbook, platform: Dict[str, Any]) -> None:
    """
    Xây sheet chiến lược cho 1 nền tảng. Có 2 style pillar table:

    - "facebook_style": bảng dọc, Content Pillar × (Tỉ lệ, Angle, Mục tiêu, Định dạng ảnh/video)
    - "tiktok_style":   bảng ngang, Pillars thành cột, thuộc tính thành dòng

    platform = {
        "name": "FACEBOOK" | "TIKTOK" | "INSTAGRAM" | ...,
        "theme": "blue_1" | "blue_2",  # blue_1=FB theme, blue_2=TikTok theme
        "channel_title": (optional) "KÊNH TIKTOK ..." - nếu có sẽ thêm row 1 title
        "info_section": [{"label": "Mục tiêu", "value": "..."}, ...],
        "pillar_table": {
            "style": "facebook_style" | "tiktok_style",
            "title": "Content Pillar ...",
            "pillars": [...]  # structure khác nhau theo style
        }
    }

    pillar_table.pillars (facebook_style):
        [
            {
                "name": "(1) Tiệc gia đình ...",
                "ratio": 0.15,
                "goal": "Khẳng định ...",
                "angles": [
                    {
                        "angle": "Câu chuyện bữa tối ...",
                        "formats": ["photo_real", "video"]  # subset của ["photo_real", "photo_design", "video"]
                    },
                    ...
                ]
            },
            ...
        ]

    pillar_table.pillars (tiktok_style):
        [
            {
                "name": "Tiệc hải sản tại Biển Đông",
                "goal": "Lấp bàn tiệc ...",
                "audience": "Lạnh" | "Ấm/Nóng" | ...,
                "purpose": "Cho khách thấy ...",
                "ratio": 0.15,
                "angle": "Angle 1 – ...\nAngle 2 – ..."
            },
            ...
        ]
    """
    theme = platform.get("theme", "blue_1")
    header_bg = COLOR["dark_blue_1"] if theme == "blue_1" else COLOR["dark_blue_2"]
    content_bg = COLOR["light_blue_1"] if theme == "blue_1" else COLOR["light_blue_2"]

    sheet_name = platform["name"].upper()
    ws = wb.create_sheet(sheet_name)

    style = platform["pillar_table"]["style"]

    if style == "facebook_style":
        _build_platform_facebook_style(ws, platform, header_bg, content_bg)
    elif style == "tiktok_style":
        _build_platform_tiktok_style(ws, platform, header_bg, content_bg)
    else:
        raise ValueError(f"Unknown pillar_table style: {style}")


def _build_platform_facebook_style(
    ws: Worksheet,
    platform: Dict[str, Any],
    header_bg: str,
    content_bg: str,
) -> None:
    """Layout y hệt sheet FACEBOOK trong mẫu Biển Đông."""

    info = platform["info_section"]

    # ---- Info section: cột A=label, B:E=content (merge) ----
    for i, item in enumerate(info, start=1):
        a = ws.cell(row=i, column=1, value=item["label"])
        _style_header(a, header_bg, h="center")

        b = ws.cell(row=i, column=2, value=item["value"])
        _style_content(b, content_bg, h=None)
        ws.merge_cells(f"B{i}:E{i}")
        # Đảm bảo các ô merge đều có border
        for c in range(2, 6):
            _style_content(ws.cell(row=i, column=c), content_bg)

    start_row_pillar = len(info) + 2  # Cách info 1 dòng trống

    # ---- Pillar title ----
    pt = platform["pillar_table"]
    title_row = start_row_pillar
    ws.cell(row=title_row, column=3, value=pt["title"])
    ws.merge_cells(f"C{title_row}:I{title_row}")
    for c in range(3, 10):
        _style_header(ws.cell(row=title_row, column=c), header_bg)

    # ---- Header row (row title+1): Content Pillar | Tỉ lệ | Content Angle | Mục tiêu | Định dạng ----
    hdr_row = title_row + 1
    sub_row = title_row + 2
    ws.cell(row=hdr_row, column=3, value="Content Pillar")
    ws.cell(row=hdr_row, column=4, value="Tỉ lệ")
    ws.cell(row=hdr_row, column=5, value="Content Angle")
    ws.cell(row=hdr_row, column=6, value="Mục tiêu")
    ws.cell(row=hdr_row, column=7, value="Định dạng")

    # Merge: C, D, E, F across 2 rows; G:I (row hdr) for "Định dạng"
    for col in [3, 4, 5, 6]:
        ws.merge_cells(start_row=hdr_row, start_column=col, end_row=sub_row, end_column=col)
    ws.merge_cells(start_row=hdr_row, start_column=7, end_row=hdr_row, end_column=9)

    # Main headers: dark blue, bold
    for col in range(3, 10):
        _style_header(ws.cell(row=hdr_row, column=col), header_bg)

    # Sub-header row (row 9): lighter blue (3C78D8), NOT bold
    for col in range(3, 10):
        sub_cell = ws.cell(row=sub_row, column=col)
        sub_cell.font = _font(bold=False, color=COLOR["white"])
        sub_cell.fill = _fill(COLOR["medium_blue_3"])
        sub_cell.alignment = _align(h="center", v="center", wrap=True)
        sub_cell.border = BORDER_ALL

    # Sub-header text: Ảnh thật | Ảnh thiết kế | Reel/Video
    ws.cell(row=sub_row, column=7, value="Ảnh thật")
    ws.cell(row=sub_row, column=8, value="Ảnh thiết kế")
    ws.cell(row=sub_row, column=9, value="Reel/ Video")

    # ---- Pillar rows ----
    current_row = sub_row + 1
    for pillar in pt["pillars"]:
        n_angles = len(pillar["angles"])
        pillar_start = current_row
        pillar_end = current_row + n_angles - 1

        # Col C: pillar name -> medium blue fill, WHITE BOLD text (nổi bật)
        c = ws.cell(row=pillar_start, column=3, value=pillar["name"])
        c.font = _font(bold=True, color=COLOR["white"])
        c.fill = _fill(COLOR["medium_blue_2"])
        c.alignment = _align(h="center", v="center", wrap=True)
        c.border = BORDER_ALL
        if n_angles > 1:
            ws.merge_cells(start_row=pillar_start, start_column=3, end_row=pillar_end, end_column=3)
        # Đảm bảo border cho ô merge
        for r in range(pillar_start, pillar_end + 1):
            cell = ws.cell(row=r, column=3)
            cell.fill = _fill(COLOR["medium_blue_2"])
            cell.border = BORDER_ALL

        # Col D: ratio (merge), hiển thị dạng 0.15
        d = ws.cell(row=pillar_start, column=4, value=pillar["ratio"])
        _style_content(d, content_bg, h="center")
        if n_angles > 1:
            ws.merge_cells(start_row=pillar_start, start_column=4, end_row=pillar_end, end_column=4)

        # Col F: goal (merge)
        f = ws.cell(row=pillar_start, column=6, value=pillar["goal"])
        _style_content(f, content_bg, h=None)
        if n_angles > 1:
            ws.merge_cells(start_row=pillar_start, start_column=6, end_row=pillar_end, end_column=6)

        # Đảm bảo border đầy đủ cho ô merge
        for r in range(pillar_start, pillar_end + 1):
            for col in [4, 6]:
                if ws.cell(row=r, column=col).value is None:
                    _style_content(ws.cell(row=r, column=col), content_bg)

        # Col E: angle (1 row/angle) + G/H/I: format flags
        for i, angle in enumerate(pillar["angles"]):
            r = pillar_start + i
            e = ws.cell(row=r, column=5, value=angle["angle"])
            _style_content(e, content_bg, h=None)

            fmts = set(angle.get("formats", []))
            ws.cell(row=r, column=7, value=("photo_real"   in fmts))
            ws.cell(row=r, column=8, value=("photo_design" in fmts))
            ws.cell(row=r, column=9, value=("video"        in fmts))

            for col in [7, 8, 9]:
                _style_content(ws.cell(row=r, column=col), content_bg, h="center")

        current_row = pillar_end + 1

    # ---- Column widths ----
    ws.column_dimensions["A"].width = 14.75
    ws.column_dimensions["B"].width = 38.0
    ws.column_dimensions["C"].width = 18.38
    ws.column_dimensions["D"].width = 11.88
    ws.column_dimensions["E"].width = 35.5
    ws.column_dimensions["F"].width = 36.38
    for col_letter in ["G", "H", "I"]:
        ws.column_dimensions[col_letter].width = 13.5


def _build_platform_tiktok_style(
    ws: Worksheet,
    platform: Dict[str, Any],
    header_bg: str,
    content_bg: str,
) -> None:
    """Layout y hệt sheet TIKTOK trong mẫu Biển Đông."""

    info = platform["info_section"]
    channel_title = platform.get("channel_title", "")
    pt = platform["pillar_table"]
    n_pillars = len(pt["pillars"])
    # Cột pillar: B, C, D, E, ... tổng số cột = 1 (label A) + n_pillars
    last_col = 1 + n_pillars
    last_col_letter = get_column_letter(last_col)

    # Info/channel title merge width: giữ nguyên mẫu = B:C (2 cột), không phải full-width
    # Mẫu Biển Đông: info section chỉ merge B:C để text không bị trải quá rộng
    info_merge_col = 3 if n_pillars >= 2 else 2  # C nếu có >=2 pillar, ngược lại chỉ B
    info_merge_letter = get_column_letter(info_merge_col)

    start_row = 1

    # ---- Row 1 (optional channel title) ----
    if channel_title:
        c = ws.cell(row=start_row, column=2, value=channel_title)
        if info_merge_col > 2:
            ws.merge_cells(f"B{start_row}:{info_merge_letter}{start_row}")
        for col in range(2, info_merge_col + 1):
            _style_header(ws.cell(row=start_row, column=col), header_bg)
        start_row += 1

    # ---- Info section ----
    info_start = start_row
    for i, item in enumerate(info):
        r = info_start + i
        a = ws.cell(row=r, column=1, value=item["label"])
        _style_header(a, header_bg, h=None)

        b = ws.cell(row=r, column=2, value=item["value"])
        _style_content(b, content_bg, h=None)
        if info_merge_col > 2:
            ws.merge_cells(f"B{r}:{info_merge_letter}{r}")
        for col in range(2, info_merge_col + 1):
            _style_content(ws.cell(row=r, column=col), content_bg)

    pillar_title_row = info_start + len(info) + 1  # Cách info 1 dòng trống

    # ---- Pillar section title ----
    ws.cell(row=pillar_title_row, column=2, value=pt["title"])
    ws.merge_cells(f"B{pillar_title_row}:{last_col_letter}{pillar_title_row}")
    for col in range(2, last_col + 1):
        _style_header(ws.cell(row=pillar_title_row, column=col), header_bg)

    # ---- Pillar names row (row pillar_title_row + 1) ----
    names_row = pillar_title_row + 1
    for i, pillar in enumerate(pt["pillars"]):
        c = ws.cell(row=names_row, column=2 + i, value=pillar["name"])
        _style_header(c, header_bg, h="center")

    # ---- Attribute rows ----
    attrs = [
        ("Mục tiêu",          "goal"),
        ("Đối tượng khách hàng", "audience"),
        ("Mục đích",          "purpose"),
        ("Tỷ lệ ",            "ratio"),
        ("Angle",             "angle"),
    ]

    for a_idx, (label, key) in enumerate(attrs):
        r = names_row + 1 + a_idx
        a = ws.cell(row=r, column=1, value=label)
        _style_header(a, header_bg, h=None)

        for i, pillar in enumerate(pt["pillars"]):
            val = pillar.get(key, "")
            b = ws.cell(row=r, column=2 + i, value=val)
            _style_content(b, content_bg, h=None)

        # Auto-merge cho dòng "Đối tượng khách hàng": nếu nhiều pillar liên tiếp
        # có cùng value (vd "Lạnh"), merge các ô đó lại.
        # (giữ nguyên hành vi mẫu Biển Đông: B14:D14 merge = 3 pillar cùng "Lạnh")
        if key == "audience":
            _merge_consecutive_same_values(ws, r, pt["pillars"], key)

    # ---- Column widths ----
    ws.column_dimensions["A"].width = 19.5
    widths = [37.88, 40.5, 39.5, 39.88]
    for i in range(n_pillars):
        w = widths[i] if i < len(widths) else 38.0
        ws.column_dimensions[get_column_letter(2 + i)].width = w


# ============================================================
# SHEET 4+: CONTENT MAPPING (weekly calendar)
# ============================================================

def build_content_mapping_sheet(wb: Workbook, platform: Dict[str, Any]) -> None:
    """
    platform = {
        "name": "FACEBOOK",
        "content_mapping": {
            "style": "dual_post" | "single_post",
            "weeks": [...]
        }
    }
    Note: CONTENT MAPPING sheets luôn dùng theme blue_1 (1C4587 + C9DAF8) bất kể
    theme của sheet chiến lược — theo đúng mẫu Biển Đông.
    """
    # Luôn dùng blue_1 theme cho content mapping
    header_bg = COLOR["dark_blue_1"]
    content_bg = COLOR["light_blue_1"]

    sheet_name = f"{platform['name'].upper()} CONTENT MAPPING"
    ws = wb.create_sheet(sheet_name)

    cm = platform["content_mapping"]
    weeks = cm["weeks"]

    def _style_cm_label(cell, bold: bool = False):
        """Label ở cột A: dark blue bg, white text, bold tuỳ theo row."""
        cell.font = _font(bold=bold, color=COLOR["white"])
        cell.fill = _fill(header_bg)
        cell.alignment = _align(h="center", v="center", wrap=True)
        cell.border = BORDER_ALL

    current_row = 1
    for week in weeks:
        n_days = week.get("days_count", 7)
        n_cols = n_days + 1  # +1 for label column A
        last_col_letter = get_column_letter(n_cols)

        # ---- Row: "Week" (A1, bold) | "Tuần X" (B..last, dark blue bold white) ----
        a = ws.cell(row=current_row, column=1, value="Week")
        _style_cm_label(a, bold=True)

        b = ws.cell(row=current_row, column=2, value=week["label"])
        # "Tuần X" hiển thị như header: dark blue bg, white bold
        b.font = _font(bold=True, color=COLOR["white"])
        b.fill = _fill(header_bg)
        b.alignment = _align(h="center", v="center", wrap=True)
        b.border = BORDER_ALL
        ws.merge_cells(f"B{current_row}:{last_col_letter}{current_row}")
        for col in range(2, n_cols + 1):
            c = ws.cell(row=current_row, column=col)
            c.fill = _fill(header_bg)
            c.border = BORDER_ALL

        # ---- Row: "Date" (bold=False) | Thứ 2..Chủ nhật (light blue content) ----
        date_row = current_row + 1
        dates_row = current_row + 2
        a = ws.cell(row=date_row, column=1, value="Date")
        _style_cm_label(a, bold=False)
        ws.merge_cells(start_row=date_row, start_column=1, end_row=dates_row, end_column=1)
        _style_cm_label(ws.cell(row=dates_row, column=1), bold=False)

        day_names = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật"]
        for i in range(n_days):
            cell = ws.cell(row=date_row, column=2 + i, value=day_names[i])
            _style_content(cell, content_bg, h="center")

        # Actual dates — transparent fill (no background), date format
        dates = week.get("dates")
        if not dates:
            start = datetime.fromisoformat(week["start_date"])
            dates = [start + timedelta(days=i) for i in range(n_days)]
        else:
            dates = [datetime.fromisoformat(d) if isinstance(d, str) else d for d in dates]

        for i, d in enumerate(dates):
            cell = ws.cell(row=dates_row, column=2 + i, value=d)
            cell.number_format = "yyyy-mm-dd"
            _style_plain(cell, h="center")

        # ---- Slot rows ----
        next_row = dates_row + 1
        for slot in week["slots"]:
            # Row 1 of slot: "Pillar" (bold=False) | 7 pillar names (no fill)
            a = ws.cell(row=next_row, column=1, value="Pillar")
            _style_cm_label(a, bold=False)
            for i in range(n_days):
                item = slot["items"][i] if i < len(slot["items"]) else {"pillar": "", "title": ""}
                cell = ws.cell(row=next_row, column=2 + i, value=item["pillar"])
                _style_plain(cell, h="center")

            # Row 2 of slot: "Title Video"/"Title Bài viết"/"Title" | 7 titles
            t_row = next_row + 1
            a = ws.cell(row=t_row, column=1, value=slot["slot_name"])
            _style_cm_label(a, bold=False)
            for i in range(n_days):
                item = slot["items"][i] if i < len(slot["items"]) else {"pillar": "", "title": ""}
                cell = ws.cell(row=t_row, column=2 + i, value=item["title"])
                _style_plain(cell, h=None)

            next_row = t_row + 1

        current_row = next_row  # Không có gap giữa weeks (theo mẫu)

    # ---- Column widths ----
    ws.column_dimensions["A"].width = 14.0
    for i in range(1, 8):
        ws.column_dimensions[get_column_letter(1 + i)].width = 27.5


# ============================================================
# MAIN ENTRY
# ============================================================

def build_plan(data: Dict[str, Any], output_path: str) -> str:
    """
    Xây file Excel Plan Content Marketing theo mẫu.

    data structure:
    {
        "company_name": "Tên doanh nghiệp",
        "customer_portrait": {
            "segments": [...],
            "dimensions": [{"label": "...", "values": [...]}, ...]
        },
        "general_direction": [{"label": "...", "value": "..."}, ...],
        "platforms": [
            {
                "name": "FACEBOOK",
                "theme": "blue_1",
                "info_section": [...],
                "pillar_table": {...},
                "content_mapping": {...}
            },
            ...
        ]
    }
    """
    wb = Workbook()
    # Remove default sheet
    default = wb.active
    wb.remove(default)

    # 1. Customer portrait
    build_customer_portrait_sheet(wb, data)

    # 2. General direction
    build_general_direction_sheet(wb, data)

    # 3+. Per-platform: strategy sheet + content mapping sheet
    for platform in data.get("platforms", []):
        build_platform_sheet(wb, platform)
        if "content_mapping" in platform and platform["content_mapping"].get("weeks"):
            build_content_mapping_sheet(wb, platform)

    wb.save(output_path)
    return output_path


# ============================================================
# CLI
# ============================================================

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        print("Usage: python build_plan.py <input.json> <output.xlsx>")
        sys.exit(1)

    in_path, out_path = sys.argv[1], sys.argv[2]
    with open(in_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    build_plan(data, out_path)
    print(f"✅ Đã tạo: {out_path}")
